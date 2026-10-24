import base64
import csv
import io
import logging
from collections import OrderedDict

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools.float_utils import float_compare

_logger = logging.getLogger(__name__)

# accepted header names (lower case) -> internal key
COLUMNS = {
    'product': 'product', 'finished product': 'product', 'bom product': 'product',
    'reference': 'reference', 'bom reference': 'reference', 'code': 'reference',
    'bom quantity': 'bom_qty', 'bom qty': 'bom_qty', 'quantity produced': 'bom_qty',
    'bom uom': 'bom_uom', 'bom unit': 'bom_uom',
    'bom type': 'bom_type', 'type': 'bom_type',
    'component': 'component', 'component product': 'component',
    'component quantity': 'component_qty', 'component qty': 'component_qty', 'quantity': 'component_qty',
    'component uom': 'component_uom', 'component unit': 'component_uom', 'uom': 'component_uom', 'unit': 'component_uom',
    'operation': 'operation', 'consumed in operation': 'operation',
    'work center': 'workcenter', 'workcenter': 'workcenter',
    'duration': 'duration', 'duration (min)': 'duration', 'operation duration': 'duration',
}
REQUIRED = ('product', 'component', 'component_qty')
BOM_TYPES = {'normal': 'normal', 'manufacture': 'normal', 'manufacture this product': 'normal',
             'kit': 'phantom', 'phantom': 'phantom'}
SAMPLE_ROWS = [
    ['Product', 'Reference', 'BoM Quantity', 'BoM Type', 'Component', 'Component Quantity', 'Component UoM',
     'Operation', 'Work Center', 'Duration (min)'],
    ['FURN_8888', 'TABLE-V1', '1', 'normal', 'FURN_0001', '4', 'Units', 'Assembly', 'Assembly line', '30'],
    ['FURN_8888', 'TABLE-V1', '1', 'normal', 'FURN_0002', '1', 'Units', 'Assembly', 'Assembly line', '30'],
    ['FURN_8888', 'TABLE-V1', '1', 'normal', 'FURN_0003', '16', 'Units', 'Assembly', 'Assembly line', '30'],
    ['FURN_9999', 'DESK-KIT', '1', 'kit', 'FURN_0004', '2', 'Units', '', '', ''],
]


class MrpBomImportWizard(models.TransientModel):
    _name = 'mrp.bom.import.wizard'
    _description = 'Import Bills of Materials'

    state = fields.Selection([('draft', 'Upload'), ('done', 'Result')], default='draft')
    file = fields.Binary("File", required=True, attachment=False)
    filename = fields.Char()
    product_match = fields.Selection(
        [('default_code', 'Internal reference'), ('barcode', 'Barcode'), ('name', 'Name')],
        string="Find products by", default='default_code', required=True)
    import_mode = fields.Selection(
        [('create', 'Always create new BoMs'),
         ('update', 'Update the BoM with the same product and reference (replace its lines)'),
         ('skip', 'Skip products that already have a BoM with this reference')],
        default='update', required=True)
    on_error = fields.Selection(
        [('skip', 'Import valid BoMs, report the rows in error'),
         ('stop', 'Import nothing if any row is in error')],
        default='skip', required=True)
    create_missing_components = fields.Boolean(
        help="Create a storable product for components that cannot be found (name and reference from the file).")
    default_bom_type = fields.Selection([('normal', 'Manufacture this product'), ('phantom', 'Kit')],
                                        default='normal', required=True)
    delimiter = fields.Selection([('auto', 'Detect'), (',', 'Comma'), (';', 'Semicolon'), ('\t', 'Tab')],
                                 string="CSV separator", default='auto', required=True)
    company_id = fields.Many2one('res.company', default=lambda self: self.env.company, required=True)

    created_count = fields.Integer(readonly=True)
    updated_count = fields.Integer(readonly=True)
    skipped_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)
    bom_ids = fields.Many2many('mrp.bom', string="Imported BoMs", readonly=True)
    result_line_ids = fields.One2many('mrp.bom.import.result', 'wizard_id', readonly=True)

    # ------------------------------------------------------------------
    # Reading the file
    # ------------------------------------------------------------------
    def _read_rows(self):
        """Return [(row_number, {key: value})] with normalized keys."""
        self.ensure_one()
        data = self._file_content()
        name = (self.filename or '').lower()
        if name.endswith(('.xlsx', '.xlsm')):
            table = self._read_xlsx(data)
        elif name.endswith('.csv') or not name:
            table = self._read_csv(data)
        else:
            raise UserError(_("Unsupported file type: use a .csv or .xlsx file."))
        if not table:
            raise UserError(_("The file is empty."))
        header = [str(h or '').strip().lower() for h in table[0]]
        keys = [COLUMNS.get(h) for h in header]
        missing = [k for k in REQUIRED if k not in keys]
        if missing:
            labels = {'product': 'Product', 'component': 'Component', 'component_qty': 'Component Quantity'}
            raise UserError(_("Missing column(s): %(cols)s.\nFound: %(found)s.\nDownload the sample file to see the expected layout.",
                              cols=', '.join(labels[m] for m in missing), found=', '.join(h for h in header if h)))
        rows = []
        for number, values in enumerate(table[1:], start=2):
            row = {key: (str(value).strip() if value is not None else '') for key, value in zip(keys, values) if key}
            if any(row.values()):
                rows.append((number, row))
        return rows

    def _file_content(self):
        value = self.file
        if isinstance(value, (bytes, str)):  # base64 (Odoo 19)
            return base64.b64decode(value)
        return bytes(value)  # BinaryValue (Odoo 20)

    @staticmethod
    def _decode(data):
        """UTF-8 (with or without BOM) first; else Windows-1252 / Latin-1 as saved by Excel in Europe."""
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode('latin-1')

    def _read_csv(self, data):
        text = self._decode(data)
        delimiter = self.delimiter
        if delimiter == 'auto':
            try:
                delimiter = csv.Sniffer().sniff(text.splitlines()[0] if text else ',', delimiters=',;\t').delimiter
            except csv.Error:
                delimiter = ','
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))

    def _read_xlsx(self, data):
        try:
            import openpyxl  # noqa: PLC0415 - shipped with Odoo
        except ImportError as e:
            raise UserError(_("The Python library openpyxl is required to read .xlsx files.")) from e
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 - zip, XML or format errors of a damaged/renamed file
            raise UserError(_("This file cannot be read as an Excel .xlsx workbook (%s). "
                              "Save it again as .xlsx or export it as .csv.", e)) from e
        sheet = workbook.worksheets[0]
        table = []
        for values in sheet.iter_rows(values_only=True):
            # 4.0 -> 4 so codes typed as numbers (barcodes, references) still match
            table.append([int(v) if isinstance(v, float) and v.is_integer() else v for v in values])
        return table

    # ------------------------------------------------------------------
    # Matching
    # ------------------------------------------------------------------
    def _find_product(self, value, cache):
        if value in cache:
            return cache[value]
        Product = self.env['product.product'].with_context(active_test=True)
        field = self.product_match
        product = Product.search([(field, '=', value)], limit=2)
        if not product and field != 'name':
            product = Product.search([('name', '=', value)], limit=2)
        if len(product) > 1:
            raise ValueError(_("Several products match %s.", value))
        cache[value] = product
        return product

    def _find_uom(self, value, product, cache):
        if not value:
            return product.uom_id
        key = value.lower()
        if key not in cache:
            uom = self.env['uom.uom'].search([('name', '=ilike', value)], limit=1)
            cache[key] = uom
        uom = cache[key]
        if not uom:
            raise ValueError(_("Unknown unit of measure %s.", value))
        if not uom._has_common_reference(product.uom_id):
            raise ValueError(_("Unit %(uom)s cannot be used for %(product)s.", uom=value, product=product.display_name))
        return uom

    def _find_workcenter(self, value, cache):
        if value not in cache:
            workcenter = self.env['mrp.workcenter'].search([('name', '=ilike', value)], limit=1)
            if not workcenter:
                workcenter = self.env['mrp.workcenter'].create({'name': value, 'company_id': self.company_id.id})
            cache[value] = workcenter
        return cache[value]

    def _number(self, value, label):
        try:
            number = float(str(value).replace(',', '.')) if value not in ('', None) else None
        except ValueError as e:
            raise ValueError(_("%(label)s is not a number: %(value)s", label=label, value=value)) from e
        return number

    # ------------------------------------------------------------------
    # Import
    # ------------------------------------------------------------------
    def _group_rows(self, rows):
        """Rows sharing (product, reference) form one BoM, in file order."""
        groups = OrderedDict()
        for number, row in rows:
            groups.setdefault((row.get('product', ''), row.get('reference', '')), []).append((number, row))
        return groups

    def _prepare_bom(self, key, rows, caches):
        """Validate one BoM; return (values, operations, lines) or raise ValueError with the row number."""
        product_value, reference = key
        first_number, first = rows[0]
        if not product_value:
            raise ValueError((first_number, _("Product is empty.")))
        product = self._find_product(product_value, caches['product'])
        if not product:
            raise ValueError((first_number, _("Product %s not found.", product_value)))
        try:
            bom_qty = self._number(first.get('bom_qty'), _("BoM Quantity")) or 1.0
            bom_uom = self._find_uom(first.get('bom_uom'), product, caches['uom'])
        except ValueError as e:
            raise ValueError((first_number, str(e.args[0]))) from e
        bom_type = BOM_TYPES.get((first.get('bom_type') or '').lower(), self.default_bom_type) \
            if first.get('bom_type') else self.default_bom_type
        if first.get('bom_type') and (first['bom_type'].lower() not in BOM_TYPES):
            raise ValueError((first_number, _("Unknown BoM type %s (use normal or kit).", first['bom_type'])))
        operations = OrderedDict()
        lines = []
        for number, row in rows:
            try:
                if not row.get('component'):
                    raise ValueError(_("Component is empty."))
                component = self._find_product(row['component'], caches['product'])
                if not component:
                    if not self.create_missing_components:
                        raise ValueError(_("Component %s not found.", row['component']))
                    component = self._create_component(row['component'])
                    caches['product'][row['component']] = component
                if component == product:
                    raise ValueError(_("A product cannot be its own component."))
                qty = self._number(row.get('component_qty'), _("Component Quantity"))
                if qty is None or float_compare(qty, 0.0, precision_digits=6) <= 0:
                    raise ValueError(_("Component Quantity must be greater than 0."))
                uom = self._find_uom(row.get('component_uom'), component, caches['uom'])
                operation = row.get('operation') or ''
                if operation and operation not in operations:
                    if not row.get('workcenter'):
                        raise ValueError(_("Operation %s needs a Work Center.", operation))
                    operations[operation] = {
                        'name': operation,
                        'workcenter_id': self._find_workcenter(row['workcenter'], caches['workcenter']).id,
                        'time_cycle_manual': self._number(row.get('duration'), _("Duration")) or 0.0,
                        'time_mode': 'manual',
                    }
            except ValueError as e:
                raise ValueError((number, str(e.args[0]))) from e
            lines.append({'product_id': component.id, 'product_qty': qty, 'uom': uom.id, 'operation': operation})
        variant_specific = len(product.product_tmpl_id.product_variant_ids) > 1
        values = {
            'product_tmpl_id': product.product_tmpl_id.id,
            'product_id': product.id if variant_specific else False,
            'code': reference or False,
            'product_qty': bom_qty,
            self._uom_field('mrp.bom'): bom_uom.id,
            'type': bom_type,
            'company_id': self.company_id.id,
        }
        return values, operations, lines

    def _uom_field(self, model):
        """The unit field of BoMs and BoM lines: product_uom_id up to Odoo 19, uom_id from Odoo 20."""
        return 'uom_id' if 'uom_id' in self.env[model]._fields else 'product_uom_id'

    def _create_component(self, value):
        values = {'name': value, 'type': 'consu', 'is_storable': True}
        if self.product_match == 'default_code':
            values['default_code'] = value
        elif self.product_match == 'barcode':
            values['barcode'] = value
        return self.env['product.product'].create(values)

    def _existing_bom(self, values):
        domain = [('product_tmpl_id', '=', values['product_tmpl_id']), ('code', '=', values['code']),
                  ('company_id', 'in', [values['company_id'], False])]
        if values['product_id']:
            domain.append(('product_id', '=', values['product_id']))
        return self.env['mrp.bom'].search(domain, limit=1)

    def _write_bom(self, values, operations, lines):
        """Create or update the BoM. Returns (bom, 'created' | 'updated' | 'skipped')."""
        existing = self._existing_bom(values) if self.import_mode != 'create' else self.env['mrp.bom']
        if existing and self.import_mode == 'skip':
            return existing, 'skipped'
        if existing:
            existing.bom_line_ids.unlink()
            existing.operation_ids.unlink()
            existing.write(values)
            bom, status = existing, 'updated'
        else:
            bom, status = self.env['mrp.bom'].create(values), 'created'
        op_records = {}
        for name, op_values in operations.items():
            op_records[name] = self.env['mrp.routing.workcenter'].create(dict(op_values, bom_id=bom.id))
        uom_field = self._uom_field('mrp.bom.line')
        self.env['mrp.bom.line'].create([{
            'bom_id': bom.id, 'product_id': line['product_id'], 'product_qty': line['product_qty'],
            uom_field: line['uom'],
            'operation_id': op_records[line['operation']].id if line['operation'] else False,
        } for line in lines])
        return bom, status

    def action_import(self):
        self.ensure_one()
        rows = self._read_rows()
        groups = self._group_rows(rows)
        caches = {'product': {}, 'uom': {}, 'workcenter': {}}
        prepared, errors = [], []
        for key, group_rows in groups.items():
            try:
                prepared.append((key, group_rows, self._prepare_bom(key, group_rows, caches)))
            except ValueError as e:
                number, message = e.args[0]
                errors.append({'row': number, 'product': key[0], 'reference': key[1], 'message': message,
                               'level': 'error'})
        if errors and self.on_error == 'stop':
            prepared = []
        counts = {'created': 0, 'updated': 0, 'skipped': 0}
        boms = self.env['mrp.bom']
        results = []
        for key, group_rows, (values, operations, lines) in prepared:
            try:
                with self.env.cr.savepoint():
                    bom, status = self._write_bom(values, operations, lines)
            except Exception as e:  # noqa: BLE001 - report any database/validation error on its row
                errors.append({'row': group_rows[0][0], 'product': key[0], 'reference': key[1],
                               'message': str(e), 'level': 'error'})
                continue
            counts[status] += 1
            boms |= bom
            results.append({'row': group_rows[0][0], 'product': key[0], 'reference': key[1], 'level': status,
                            'message': _("%(count)s component(s)", count=len(lines)), 'bom_id': bom.id})
        self.write({
            'state': 'done',
            'created_count': counts['created'], 'updated_count': counts['updated'],
            'skipped_count': counts['skipped'], 'error_count': len(errors),
            'bom_ids': [(6, 0, boms.ids)],
            'result_line_ids': [(5, 0, 0)] + [(0, 0, v) for v in sorted(errors + results, key=lambda r: r['row'])],
        })
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new', 'name': _("Import result")}

    def action_open_boms(self):
        action = self.env['ir.actions.act_window']._for_xml_id('mrp.mrp_bom_form_action')
        action['domain'] = [('id', 'in', self.bom_ids.ids)]
        action['context'] = {}
        return action

    @api.model
    def action_download_sample(self, file_type='xlsx'):
        if file_type == 'csv':
            buffer = io.StringIO()
            csv.writer(buffer).writerows(SAMPLE_ROWS)
            content, name, mimetype = buffer.getvalue().encode(), 'bom_import_sample.csv', 'text/csv'
        else:
            import openpyxl  # noqa: PLC0415
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'Bills of Materials'
            sheet.append(SAMPLE_ROWS[0])
            for row in SAMPLE_ROWS[1:]:
                # quantities and duration as numbers, codes as text
                sheet.append([float(v) if i in (2, 5, 9) and v else v for i, v in enumerate(row)])
            buffer = io.BytesIO()
            workbook.save(buffer)
            content = buffer.getvalue()
            name = 'bom_import_sample.xlsx'
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        attachment = self.env['ir.attachment'].create({
            'name': name, 'raw': content, 'mimetype': mimetype,
            'res_model': self._name, 'public': False,
        })
        return {'type': 'ir.actions.act_url', 'url': f'/web/content/{attachment.id}?download=true', 'target': 'download'}

    def action_download_sample_csv(self):
        return self.action_download_sample('csv')

    def action_download_sample_xlsx(self):
        return self.action_download_sample('xlsx')

    def action_back(self):
        self.write({'state': 'draft', 'result_line_ids': [(5, 0, 0)]})
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new'}


class MrpBomImportResult(models.TransientModel):
    _name = 'mrp.bom.import.result'
    _description = 'BoM import result line'
    _order = 'row, id'

    wizard_id = fields.Many2one('mrp.bom.import.wizard', required=True, ondelete='cascade')
    row = fields.Integer("Row")
    product = fields.Char()
    reference = fields.Char()
    level = fields.Selection([('created', 'Created'), ('updated', 'Updated'), ('skipped', 'Skipped'),
                              ('error', 'Error')], string="Result")
    message = fields.Char()
    bom_id = fields.Many2one('mrp.bom', string="BoM")
